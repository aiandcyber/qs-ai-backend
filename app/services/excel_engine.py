"""Excel claim / BQ / certificate ingest and valuation export via openpyxl."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

HEADERS = [
    "Item No",
    "Description",
    "Unit",
    "Contract Qty",
    "Rate",
    "Contract Amount",
    "Previous Qty",
    "Previous Amount",
    "Claimed Qty",
    "Claimed Amount",
    "Assessed Qty",
    "Assessed Amount",
    "Difference",
    "Risk",
    "Issues",
    "Category",
]


def _f(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def read_lines(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    index = {name: i for i, name in enumerate(header)}

    def col(*names: str, default: Any = None) -> Any:
        for name in names:
            if name in index:
                return index[name]
        return default

    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        item_no = str(row[col("Item No", "Item")] or "").strip()
        if not item_no or item_no.lower().startswith("total"):
            continue
        contract_qty = _f(row[col("Contract Qty", "Qty")] if col("Contract Qty", "Qty") is not None else 0)
        rate = _f(row[col("Rate")] if col("Rate") is not None else 0)
        contract_amount = _f(
            row[col("Contract Amount", "Amount")] if col("Contract Amount", "Amount") is not None else contract_qty * rate
        )
        previous_qty = _f(row[col("Previous Qty")] if col("Previous Qty") is not None else 0)
        previous_amount = _f(
            row[col("Previous Amount")] if col("Previous Amount") is not None else previous_qty * rate
        )
        claimed_qty = _f(row[col("Claimed Qty", "This Period Qty")] if col("Claimed Qty", "This Period Qty") is not None else 0)
        claimed_amount = _f(
            row[col("Claimed Amount", "This Period Amount")]
            if col("Claimed Amount", "This Period Amount") is not None
            else claimed_qty * rate
        )
        description = str(row[col("Description")] if col("Description") is not None else "")
        unit = str(row[col("Unit")] if col("Unit") is not None else "")
        category = str(row[col("Category")] if col("Category") is not None else "works")
        out.append(
            {
                "item_no": item_no,
                "description": description,
                "unit": unit,
                "contract_qty": contract_qty,
                "rate": rate,
                "contract_amount": contract_amount,
                "previous_qty": previous_qty,
                "previous_amount": previous_amount,
                "claimed_qty": claimed_qty,
                "claimed_amount": claimed_amount,
                "category": category,
            }
        )
    return out


def write_valuation(path: Path, lines: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Valuation"
    ws.append(HEADERS)
    for line in lines:
        ws.append(
            [
                line.get("item_no"),
                line.get("description"),
                line.get("unit"),
                line.get("contract_qty"),
                line.get("rate"),
                line.get("contract_amount"),
                line.get("previous_qty"),
                line.get("previous_amount"),
                line.get("claimed_qty"),
                line.get("claimed_amount"),
                line.get("assessed_qty"),
                line.get("assessed_amount"),
                line.get("difference_amount"),
                line.get("risk"),
                "; ".join(line.get("issues") or []),
                line.get("category"),
            ]
        )
    ws2 = wb.create_sheet("Summary")
    for key, value in summary.items():
        ws2.append([key, value])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
